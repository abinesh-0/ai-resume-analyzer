import io
import logging
import os
import re
import secrets
import uuid
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path

import joblib
import mysql.connector
import pymupdf
import pytesseract
from docx import Document
from dotenv import load_dotenv
from flask import (
    Flask,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    url_for,
)
from flask_mail import Mail, Message
from mysql.connector import Error
from PIL import Image
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


# =========================================================
# ENVIRONMENT
# =========================================================

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent

UPLOAD_FOLDER = Path(
    os.getenv(
        "UPLOAD_FOLDER",
        str(BASE_DIR / "uploads")
    )
)

MODEL_PATH = Path(
    os.getenv(
        "MODEL_PATH",
        str(BASE_DIR / "models" / "career_model.pkl")
    )
)

VECTORIZER_PATH = Path(
    os.getenv(
        "VECTORIZER_PATH",
        str(BASE_DIR / "models" / "tfidf_vectorizer.pkl")
    )
)

UPLOAD_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)


# =========================================================
# FLASK APP
# =========================================================

app = Flask(__name__)

app.config.update(
    SECRET_KEY=os.getenv(
        "SECRET_KEY",
        "dev-only-change-me"
    ),

    MAX_CONTENT_LENGTH=10 * 1024 * 1024,

    OTP_EXPIRY_SECONDS=90,

    OTP_RESEND_COOLDOWN_SECONDS=30,

    OTP_MAX_ATTEMPTS=5,

    MAIL_SERVER=os.getenv(
        "MAIL_SERVER",
        "smtp.gmail.com"
    ),

    MAIL_PORT=int(
        os.getenv(
            "MAIL_PORT",
            "587"
        )
    ),

    MAIL_USE_TLS=os.getenv(
        "MAIL_USE_TLS",
        "true"
    ).lower() == "true",

    MAIL_USERNAME=os.getenv(
        "MAIL_USERNAME"
    ),

    MAIL_PASSWORD=os.getenv(
        "MAIL_PASSWORD"
    ),
)


# =========================================================
# LOGGING
# =========================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

logger = logging.getLogger(__name__)


# =========================================================
# MAIL
# =========================================================

mail = Mail(app)


# =========================================================
# FILE SETTINGS
# =========================================================

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx"
}


# =========================================================
# EMAIL VALIDATION
# =========================================================

EMAIL_RE = re.compile(
    r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$"
)

GMAIL_RE = re.compile(
    r"^[A-Za-z0-9._%+-]+@gmail\.com$",
    re.IGNORECASE
)


# =========================================================
# CAREERS
# =========================================================

CAREERS = {

    "Software Engineer": [
        "python",
        "java",
        "c++",
        "sql",
        "git",
        "data structures",
        "algorithms"
    ],

    "Full Stack Developer": [
        "html",
        "css",
        "javascript",
        "react",
        "python",
        "flask",
        "sql",
        "git",
        "rest api"
    ],

    "Frontend Developer": [
        "html",
        "css",
        "javascript",
        "react",
        "typescript",
        "git",
        "bootstrap"
    ],

    "Backend Developer": [
        "python",
        "java",
        "sql",
        "flask",
        "django",
        "fastapi",
        "rest api",
        "git",
        "docker"
    ],

    "Python Developer": [
        "python",
        "sql",
        "git",
        "flask",
        "django",
        "fastapi",
        "rest api",
        "docker"
    ],

    "Java Developer": [
        "java",
        "sql",
        "spring",
        "rest api",
        "git",
        "docker"
    ],

    "Web Developer": [
        "html",
        "css",
        "javascript",
        "bootstrap",
        "git",
        "sql"
    ],

    "Mobile App Developer": [
        "java",
        "kotlin",
        "flutter",
        "dart",
        "firebase",
        "git"
    ],

    "AI Engineer": [
        "python",
        "numpy",
        "pandas",
        "machine learning",
        "deep learning",
        "tensorflow",
        "pytorch",
        "git"
    ],

    "Machine Learning Engineer": [
        "python",
        "numpy",
        "pandas",
        "scikit-learn",
        "machine learning",
        "deep learning",
        "tensorflow",
        "pytorch",
        "git"
    ],

    "AI/ML Engineer": [
        "python",
        "numpy",
        "pandas",
        "machine learning",
        "deep learning",
        "tensorflow",
        "pytorch",
        "opencv",
        "git"
    ],

    "Data Scientist": [
        "python",
        "sql",
        "pandas",
        "numpy",
        "machine learning",
        "data analysis",
        "statistics",
        "matplotlib",
        "scikit-learn"
    ],

    "Data Analyst": [
        "python",
        "sql",
        "excel",
        "pandas",
        "numpy",
        "power bi",
        "tableau",
        "data analysis",
        "statistics"
    ],

    "NLP Engineer": [
        "python",
        "numpy",
        "pandas",
        "machine learning",
        "deep learning",
        "nlp",
        "tensorflow",
        "pytorch"
    ],

    "Computer Vision Engineer": [
        "python",
        "numpy",
        "opencv",
        "computer vision",
        "machine learning",
        "deep learning",
        "pytorch",
        "tensorflow"
    ],

    "Cloud Engineer": [
        "linux",
        "aws",
        "azure",
        "docker",
        "kubernetes",
        "terraform",
        "git",
        "networking"
    ],

    "DevOps Engineer": [
        "linux",
        "git",
        "docker",
        "kubernetes",
        "aws",
        "azure",
        "ci/cd",
        "terraform"
    ],

    "Cloud Architect": [
        "aws",
        "azure",
        "cloud architecture",
        "networking",
        "linux",
        "docker",
        "kubernetes",
        "terraform"
    ],

    "Cybersecurity Engineer": [
        "networking",
        "linux",
        "python",
        "cybersecurity",
        "cryptography",
        "ethical hacking",
        "siem",
        "git"
    ],

    "Security Analyst": [
        "networking",
        "linux",
        "cybersecurity",
        "siem",
        "incident response",
        "python",
        "cryptography"
    ],

    'Data Engineer': ['python', 'sql', 'data modeling', 'etl', 'data warehouse', 'spark', 'kafka', 'airflow', 'cloud', 'git'],
    'MLOps Engineer': ['python', 'machine learning', 'docker', 'kubernetes', 'ci/cd', 'cloud', 'mlops', 'terraform', 'git', 'monitoring'],
    'Site Reliability Engineer': ['linux', 'networking', 'python', 'bash', 'cloud', 'docker', 'kubernetes', 'ci/cd', 'terraform', 'monitoring', 'incident response'],
    'QA Engineer': ['software testing', 'test automation', 'sql', 'api testing', 'selenium', 'pytest', 'bug tracking', 'ci/cd', 'git'],
    'Automation Test Engineer': ['software testing', 'python', 'selenium', 'pytest', 'api testing', 'test automation', 'sql', 'jenkins', 'ci/cd', 'git'],
    'Network Engineer': ['networking', 'tcp/ip', 'ip addressing', 'routing', 'switching', 'dns', 'dhcp', 'firewall', 'vpn', 'network monitoring'],
    'Systems Engineer': ['operating systems', 'linux', 'windows server', 'networking', 'bash', 'python', 'virtualization', 'cloud', 'monitoring', 'automation'],
    'Database Administrator': ['sql', 'mysql', 'postgresql', 'database architecture', 'backup and recovery', 'database security', 'indexes', 'performance tuning', 'high availability', 'linux'],
    'Solutions Architect': ['system design', 'cloud', 'networking', 'security', 'rest api', 'microservices', 'databases', 'docker', 'scalability', 'architecture'],
    'Generative AI Engineer': ['python', 'machine learning', 'deep learning', 'transformers', 'llms', 'prompt engineering', 'rag', 'vector database', 'llm evaluation', 'api', 'deployment'],
    'AI Application Engineer': ['python', 'rest api', 'llms', 'prompt engineering', 'rag', 'vector database', 'fastapi', 'flask', 'sql', 'deployment'],
    'Robotics Engineer': ['python', 'c++', 'robotics', 'sensors', 'control systems', 'ros', 'computer vision', 'robot simulation', 'navigation', 'embedded systems'],
    'Embedded Systems Engineer': ['c', 'c++', 'microcontrollers', 'digital electronics', 'interrupts', 'rtos', 'device drivers', 'debugging', 'embedded systems', 'uart', 'spi', 'i2c'],
    'Firmware Engineer': ['c', 'c++', 'microcontrollers', 'firmware', 'device drivers', 'rtos', 'communication protocols', 'debugging', 'testing', 'embedded systems'],
    'IoT Engineer': ['python', 'c++', 'iot', 'sensors', 'arduino', 'raspberry pi', 'mqtt', 'networking', 'cloud', 'data processing', 'iot security'],
    'Hardware Engineer': ['circuit design', 'analog electronics', 'digital electronics', 'microcontrollers', 'pcb design', 'verilog', 'hdl', 'testing', 'debugging', 'embedded systems'],
    'Blockchain Developer': ['javascript', 'blockchain', 'ethereum', 'solidity', 'smart contracts', 'web3', 'wallet integration', 'cryptography', 'security', 'dapp'],
    'Game Developer': ['c++', 'c#', 'game engine', '2d/3d math', 'game physics', 'gameplay programming', 'unity', 'unreal engine', '3d assets', 'optimization'],
    'AR/VR Developer': ['c#', 'unity', '3d math', '3d graphics', 'ar/vr', 'interaction design', 'computer vision', 'blender', 'optimization', 'spatial computing'],
    'Platform Engineer': ['linux', 'git', 'docker', 'kubernetes', 'cloud', 'terraform', 'ci/cd', 'developer platform', 'observability', 'security', 'python'],
}


# =========================================================
# SKILLS
# =========================================================

SKILLS = sorted(
    {
        skill
        for skills in CAREERS.values()
        for skill in skills
    }
    |
    {
        "javascript",
        "typescript",
        "html",
        "css",
        "java",
        "python",
        "c++",
        "c",
        "sql",
        "mysql",
        "sqlite",
        "flask",
        "django",
        "fastapi",
        "react",
        "node.js",
        "nodejs",
        "spring",
        "kotlin",
        "flutter",
        "dart",
        "firebase",
        "machine learning",
        "deep learning",
        "artificial intelligence",
        "data science",
        "data analysis",
        "pandas",
        "numpy",
        "scikit-learn",
        "tensorflow",
        "pytorch",
        "opencv",
        "computer vision",
        "nlp",
        "statistics",
        "matplotlib",
        "excel",
        "power bi",
        "tableau",
        "git",
        "github",
        "docker",
        "kubernetes",
        "aws",
        "azure",
        "terraform",
        "linux",
        "networking",
        "rest api",
        "bootstrap",
        "tailwind",
        "mongodb",
        "data structures",
        "algorithms",
        "cryptography",
        "cybersecurity",
        "ethical hacking",
        "siem",
        "incident response",
        "ci/cd",
        "cloud architecture",
    },
    key=len,
    reverse=True
)


# =========================================================
# SKILL ALIASES
# =========================================================

SKILL_ALIASES = {

    "nodejs": "node.js",

    "node js": "node.js",

    "restful api": "rest api",

    "rest-api": "rest api",

    "powerbi": "power bi",

    "scikit learn": "scikit-learn",

    "machine-learning": "machine learning",

    "deep-learning": "deep learning",

    "computer-vision": "computer vision",

    "c plus plus": "c++",
}


# =========================================================
# LEARNING TOPICS
# =========================================================

LEARNING_TOPICS = {

    "python": [
        "Syntax & data types",
        "Functions",
        "OOP",
        "Exceptions",
        "Modules",
        "Virtual environments",
        "APIs"
    ],

    "java": [
        "Core Java",
        "OOP",
        "Collections",
        "Exception handling",
        "Streams",
        "Spring basics"
    ],

    "javascript": [
        "Variables & functions",
        "DOM",
        "ES6+",
        "Async/await",
        "Fetch API",
        "Modules"
    ],

    "typescript": [
        "Types",
        "Interfaces",
        "Generics",
        "Modules",
        "Type-safe React"
    ],

    "html": [
        "Semantic HTML",
        "Forms",
        "Accessibility",
        "SEO basics"
    ],

    "css": [
        "Selectors",
        "Box model",
        "Flexbox",
        "Grid",
        "Responsive design",
        "Animations"
    ],

    "sql": [
        "SELECT",
        "JOINs",
        "GROUP BY",
        "Subqueries",
        "Indexes",
        "Transactions"
    ],

    "mysql": [
        "Schema design",
        "Queries",
        "Indexes",
        "Constraints",
        "Transactions",
        "Optimization"
    ],

    "machine learning": [
        "Data preprocessing",
        "Regression",
        "Classification",
        "Clustering",
        "Evaluation",
        "Feature engineering",
        "Tuning"
    ],

    "deep learning": [
        "Neural networks",
        "Backpropagation",
        "CNN",
        "RNN",
        "Transfer learning",
        "Model training"
    ],

    "tensorflow": [
        "Tensors",
        "Datasets",
        "Keras models",
        "Callbacks",
        "Training",
        "Saving/deployment"
    ],

    "pytorch": [
        "Tensors",
        "Datasets",
        "Autograd",
        "Neural networks",
        "Training loops",
        "Model saving"
    ],

    "pandas": [
        "Series/DataFrame",
        "Cleaning",
        "Filtering",
        "GroupBy",
        "Merge",
        "Time series"
    ],

    "numpy": [
        "Arrays",
        "Indexing",
        "Vectorization",
        "Broadcasting",
        "Linear algebra"
    ],

    "scikit-learn": [
        "Preprocessing",
        "Pipelines",
        "Models",
        "Cross-validation",
        "Metrics",
        "Tuning"
    ],

    "opencv": [
        "Image basics",
        "Filtering",
        "Thresholding",
        "Contours",
        "Object detection",
        "Video processing"
    ],

    "computer vision": [
        "Image processing",
        "Feature extraction",
        "CNNs",
        "Detection",
        "Segmentation",
        "Deployment"
    ],

    "nlp": [
        "Text cleaning",
        "Tokenization",
        "Embeddings",
        "Classification",
        "Transformers",
        "Evaluation"
    ],

    "data analysis": [
        "Cleaning",
        "EDA",
        "Statistics",
        "Visualization",
        "Insights",
        "Reporting"
    ],

    "statistics": [
        "Descriptive statistics",
        "Probability",
        "Distributions",
        "Hypothesis testing",
        "Correlation",
        "Regression"
    ],

    "git": [
        "Repository",
        "Commit",
        "Branch",
        "Merge",
        "Pull request",
        "Conflict resolution"
    ],

    "docker": [
        "Images",
        "Containers",
        "Dockerfile",
        "Volumes",
        "Networks",
        "Compose"
    ],

    "aws": [
        "IAM",
        "EC2",
        "S3",
        "RDS",
        "VPC",
        "Cloud monitoring"
    ],

    "azure": [
        "Identity",
        "VMs",
        "Storage",
        "Databases",
        "Networking",
        "Monitoring"
    ],

    "kubernetes": [
        "Pods",
        "Deployments",
        "Services",
        "ConfigMaps",
        "Secrets",
        "Scaling"
    ],

    "linux": [
        "Shell",
        "Files",
        "Permissions",
        "Processes",
        "Networking",
        "Package management"
    ],

    "networking": [
        "OSI/TCP-IP",
        "IP addressing",
        "DNS",
        "HTTP/HTTPS",
        "Routing",
        "Firewalls"
    ],

    "cybersecurity": [
        "Security fundamentals",
        "Threats",
        "Authentication",
        "Network security",
        "Logging",
        "Incident response"
    ],

    "power bi": [
        "Power Query",
        "Data modeling",
        "DAX",
        "Visuals",
        "Dashboards",
        "Publishing"
    ],

    "tableau": [
        "Connections",
        "Calculated fields",
        "Charts",
        "Dashboards",
        "Filters",
        "Publishing"
    ],

    "react": [
        "Components",
        "Props/state",
        "Hooks",
        "Routing",
        "API integration",
        "Performance"
    ],

    "flask": [
        "Routes",
        "Templates",
        "Forms",
        "Sessions",
        "Database integration",
        "Deployment"
    ],

    "django": [
        "Apps",
        "Models",
        "Views",
        "Templates",
        "Forms",
        "REST APIs"
    ],

    "fastapi": [
        "Routes",
        "Pydantic",
        "Validation",
        "Dependency injection",
        "Async APIs",
        "Deployment"
    ],

    "rest api": [
        "HTTP methods",
        "Status codes",
        "JSON",
        "Authentication",
        "Validation",
        "Error handling"
    ],

    "terraform": [
        "Providers",
        "Resources",
        "Variables",
        "State",
        "Modules",
        "Remote state"
    ],

    "kotlin": [
        "Syntax",
        "OOP",
        "Collections",
        "Coroutines",
        "Android basics"
    ],

    "flutter": [
        "Dart",
        "Widgets",
        "Layouts",
        "State management",
        "APIs",
        "Build/release"
    ],

    "dart": [
        "Syntax",
        "OOP",
        "Collections",
        "Async programming",
        "Flutter integration"
    ],

    "firebase": [
        "Authentication",
        "Firestore",
        "Storage",
        "Cloud Functions",
        "Security rules"
    ],

    "data structures": [
        "Arrays",
        "Linked lists",
        "Stacks/queues",
        "Trees",
        "Graphs",
        "Hashing"
    ],

    "algorithms": [
        "Searching",
        "Sorting",
        "Recursion",
        "Greedy",
        "Dynamic programming",
        "Graph algorithms"
    ],
}


# =========================================================
# CAREER ROADMAPS
# =========================================================

CAREER_ROADMAPS = {

    "AI/ML Engineer": [
        "Python",
        "NumPy + Pandas",
        "Statistics",
        "Machine Learning",
        "Deep Learning",
        "Computer Vision / NLP",
        "TensorFlow / PyTorch",
        "Projects",
        "Git + Deployment",
        "Job Ready"
    ],

    "Machine Learning Engineer": [
        "Python",
        "Data Structures",
        "NumPy + Pandas",
        "Statistics",
        "Machine Learning",
        "Scikit-learn",
        "Deep Learning",
        "ML Projects",
        "MLOps + Deployment",
        "Job Ready"
    ],

    "Data Scientist": [
        "Python",
        "SQL",
        "Statistics",
        "Pandas + NumPy",
        "EDA + Visualization",
        "Machine Learning",
        "Projects",
        "Storytelling",
        "Portfolio",
        "Job Ready"
    ],

    "Data Analyst": [
        "Excel",
        "SQL",
        "Statistics",
        "Power BI / Tableau",
        "Python + Pandas",
        "EDA",
        "Dashboards",
        "Business Projects",
        "Portfolio",
        "Job Ready"
    ],

    "Full Stack Developer": [
        "HTML + CSS",
        "JavaScript",
        "React",
        "Backend with Python/Node",
        "SQL",
        "REST APIs",
        "Authentication",
        "Git + GitHub",
        "Deployment",
        "Job Ready"
    ],

    "Frontend Developer": [
        "HTML",
        "CSS",
        "JavaScript",
        "Responsive UI",
        "React",
        "APIs",
        "Git",
        "Testing",
        "Deployment",
        "Portfolio"
    ],

    "Backend Developer": [
        "Programming",
        "SQL",
        "Backend Framework",
        "REST APIs",
        "Authentication",
        "Testing",
        "Docker",
        "Cloud",
        "Projects",
        "Job Ready"
    ],

    "Cloud Engineer": [
        "Linux",
        "Networking",
        "Cloud Fundamentals",
        "AWS/Azure",
        "IAM",
        "Containers",
        "Infrastructure as Code",
        "Monitoring",
        "Projects",
        "Job Ready"
    ],

    "DevOps Engineer": [
        "Linux",
        "Git",
        "CI/CD",
        "Docker",
        "Kubernetes",
        "Cloud",
        "Terraform",
        "Monitoring",
        "Projects",
        "Job Ready"
    ],

    "Cybersecurity Engineer": [
        "Networking",
        "Linux",
        "Security Fundamentals",
        "Python",
        "Web Security",
        "SIEM",
        "Incident Response",
        "Ethical Hacking",
        "Security Projects",
        "Job Ready"
    ],

    "Software Engineer": [
        "Programming",
        "Data Structures",
        "Algorithms",
        "OOP",
        "SQL",
        "Git",
        "APIs",
        "Testing",
        "Projects",
        "Job Ready"
    ],

    'Python Developer': ['Python Fundamentals', 'OOP', 'Data Structures', 'SQL', 'Flask / Django / FastAPI', 'REST APIs', 'Testing', 'Docker', 'Git + GitHub', 'Projects', 'Job Ready'],
    'Java Developer': ['Core Java', 'OOP', 'Collections', 'SQL', 'Spring Boot', 'REST APIs', 'JPA / Hibernate', 'Testing', 'Docker', 'Projects', 'Job Ready'],
    'Web Developer': ['HTML', 'CSS', 'JavaScript', 'Responsive Design', 'Bootstrap', 'Browser APIs', 'SQL', 'Git', 'Web Security', 'Projects', 'Job Ready'],
    'Mobile App Developer': ['Programming Basics', 'Kotlin / Java or Dart', 'UI Layouts', 'Flutter / Android', 'State Management', 'REST APIs', 'Firebase', 'Testing', 'App Security', 'Build + Release', 'Portfolio'],
    'AI Engineer': ['Python', 'NumPy + Pandas', 'Statistics', 'Machine Learning', 'Deep Learning', 'Model Development', 'TensorFlow / PyTorch', 'AI Applications', 'APIs + Deployment', 'Git', 'Job Ready'],
    'NLP Engineer': ['Python', 'Text Processing', 'Statistics', 'Machine Learning', 'NLP Fundamentals', 'Embeddings', 'Transformers', 'LLM Applications', 'NLP Projects', 'Deployment', 'Job Ready'],
    'Computer Vision Engineer': ['Python', 'NumPy', 'Image Processing', 'OpenCV', 'Machine Learning', 'Deep Learning', 'CNNs', 'Object Detection', 'Vision Projects', 'Deployment', 'Job Ready'],
    'Cloud Architect': ['Cloud Fundamentals', 'Networking', 'AWS + Azure', 'Identity + Security', 'High Availability', 'Scalability', 'Microservices', 'Containers', 'Infrastructure as Code', 'Architecture Case Studies', 'Job Ready'],
    'Security Analyst': ['Networking', 'Linux', 'Security Fundamentals', 'SIEM', 'Log Analysis', 'Threat Detection', 'Incident Response', 'Vulnerability Management', 'Security Projects', 'Reporting', 'Job Ready'],
    'Data Engineer': ['Python', 'SQL', 'Data Modeling', 'ETL / ELT', 'Data Warehouses', 'Spark', 'Kafka', 'Airflow', 'Cloud Data Platforms', 'Data Quality', 'Projects', 'Job Ready'],
    'MLOps Engineer': ['Python', 'ML Fundamentals', 'Docker', 'Kubernetes', 'CI/CD', 'Cloud', 'Model Serving', 'ML Pipelines', 'Terraform', 'Monitoring', 'MLOps Projects', 'Job Ready'],
    'Site Reliability Engineer': ['Linux', 'Networking', 'Python / Bash', 'Cloud', 'Containers', 'Kubernetes', 'CI/CD', 'Infrastructure as Code', 'Observability', 'Incident Response', 'SRE Projects', 'Job Ready'],
    'QA Engineer': ['Software Testing Basics', 'Test Planning', 'Test Cases', 'Bug Tracking', 'SQL Testing', 'API Testing', 'Selenium Basics', 'Pytest', 'CI/CD Testing', 'Automation Projects', 'Job Ready'],
    'Automation Test Engineer': ['Testing Fundamentals', 'Python', 'Selenium', 'Pytest', 'API Automation', 'Test Frameworks', 'SQL', 'Jenkins / CI/CD', 'Test Reporting', 'Automation Projects', 'Job Ready'],
    'Network Engineer': ['Networking Fundamentals', 'TCP/IP', 'IP Addressing', 'Routing', 'Switching', 'DNS + DHCP', 'Firewalls', 'VPNs', 'Network Monitoring', 'Troubleshooting Projects', 'Job Ready'],
    'Systems Engineer': ['Operating Systems', 'Linux', 'Windows Server', 'Networking', 'Bash + Python', 'Virtualization', 'Cloud', 'Monitoring', 'Automation', 'Systems Projects', 'Job Ready'],
    'Database Administrator': ['SQL', 'Database Architecture', 'MySQL / PostgreSQL', 'Oracle Basics', 'Backup + Recovery', 'Database Security', 'Indexes', 'Performance Tuning', 'High Availability', 'DBA Projects', 'Job Ready'],
    'Solutions Architect': ['Architecture Fundamentals', 'Cloud', 'Networking', 'Security', 'APIs', 'Microservices', 'Databases', 'Containers', 'Scalability + Reliability', 'Architecture Case Studies', 'Job Ready'],
    'Generative AI Engineer': ['Python', 'ML Fundamentals', 'Deep Learning', 'Transformers', 'LLMs', 'Prompt Engineering', 'RAG', 'Vector Databases', 'LLM Evaluation', 'AI Application Deployment', 'GenAI Projects', 'Job Ready'],
    'AI Application Engineer': ['Python', 'APIs', 'LLM Fundamentals', 'Prompt Engineering', 'RAG', 'Vector Databases', 'FastAPI / Flask', 'Databases', 'Deployment', 'AI Product Projects', 'Job Ready'],
    'Robotics Engineer': ['Python + C++', 'Math for Robotics', 'Sensors', 'Control Systems', 'ROS', 'Computer Vision', 'Robot Simulation', 'Navigation', 'Hardware Integration', 'Robotics Projects', 'Job Ready'],
    'Embedded Systems Engineer': ['C Programming', 'C++', 'Microcontrollers', 'Digital Electronics', 'Interrupts', 'RTOS', 'Device Drivers', 'Debugging', 'Embedded Communication', 'Embedded Projects', 'Job Ready'],
    'Firmware Engineer': ['C Programming', 'C++', 'Microcontrollers', 'Firmware Architecture', 'Drivers', 'RTOS', 'Communication Protocols', 'Debugging', 'Testing', 'Firmware Projects', 'Job Ready'],
    'IoT Engineer': ['Python / C++', 'IoT Fundamentals', 'Sensors', 'Arduino / Raspberry Pi', 'MQTT', 'Networking', 'Cloud IoT', 'Data Processing', 'Security', 'IoT Projects', 'Job Ready'],
    'Hardware Engineer': ['Circuit Fundamentals', 'Analog Electronics', 'Digital Electronics', 'Microcontrollers', 'PCB Design', 'Verilog / HDL', 'Testing + Debugging', 'Embedded Integration', 'Hardware Projects', 'Documentation', 'Job Ready'],
    'Blockchain Developer': ['JavaScript', 'Blockchain Fundamentals', 'Ethereum', 'Solidity', 'Smart Contracts', 'Web3 APIs', 'Wallet Integration', 'Security', 'DApp Projects', 'Testing + Deployment', 'Job Ready'],
    'Game Developer': ['Programming', 'C++ / C#', 'Game Engine', '2D/3D Math', 'Physics', 'Gameplay Systems', 'UI + Audio', '3D Assets', 'Optimization', 'Game Projects', 'Portfolio'],
    'AR/VR Developer': ['C#', 'Unity', '3D Math', '3D Graphics', 'AR/VR Fundamentals', 'Interaction Design', 'Computer Vision', 'Blender / Assets', 'Optimization', 'AR/VR Projects', 'Portfolio'],
    'Platform Engineer': ['Linux', 'Git', 'Docker', 'Kubernetes', 'Cloud', 'Terraform', 'CI/CD', 'Developer Platforms', 'Observability', 'Security', 'Platform Projects', 'Job Ready'],
}


# =========================================================
# PROJECT RECOMMENDATIONS
# =========================================================

PROJECTS = {

    "AI/ML Engineer": [
        "Resume Skill Gap Analyzer",
        "Vehicle Number Plate Recognition",
        "Medical Image Classifier",
        "Document Intelligence System"
    ],

    "Machine Learning Engineer": [
        "Customer Churn Predictor",
        "Fraud Detection System",
        "Recommendation Engine",
        "ML Model Deployment API"
    ],

    "Data Scientist": [
        "Sales Forecasting",
        "Customer Segmentation",
        "Churn Prediction",
        "End-to-End Business Analytics"
    ],

    "Data Analyst": [
        "Sales Dashboard",
        "Student Performance Dashboard",
        "Customer Analytics",
        "Business KPI Dashboard"
    ],

    "Full Stack Developer": [
        "Job Portal",
        "E-commerce Platform",
        "Student Management System",
        "Real-time Task Manager"
    ],

    "Frontend Developer": [
        "Portfolio Website",
        "Admin Dashboard",
        "E-commerce UI",
        "Interactive Analytics Dashboard"
    ],

    "Backend Developer": [
        "REST API Platform",
        "Authentication Service",
        "Inventory API",
        "Payment-ready Order API"
    ],

    "Cloud Engineer": [
        "Cloud-hosted Web App",
        "Serverless File Processor",
        "Cloud Monitoring Dashboard",
        "Highly Available API"
    ],

    "DevOps Engineer": [
        "CI/CD Pipeline",
        "Dockerized Flask App",
        "Kubernetes Deployment",
        "Infrastructure Automation"
    ],

    "Cybersecurity Engineer": [
        "Security Log Analyzer",
        "Vulnerability Lab",
        "Phishing Detection Prototype",
        "Incident Response Dashboard"
    ],

    "Software Engineer": [
        "Task Management API",
        "Student Management System",
        "E-commerce Backend",
        "Online Coding Platform"
    ],

    'Python Developer': ['Resume Analyzer API', 'Student Management API', 'Expense Tracker Backend', 'Flask E-commerce API'],
    'Java Developer': ['Spring Boot Job Portal API', 'Banking API', 'Inventory Management System', 'Employee Management Backend'],
    'Web Developer': ['Business Website', 'Blog Platform', 'Event Registration Website', 'Responsive E-commerce Site'],
    'Mobile App Developer': ['Expense Tracker App', 'College Companion App', 'Food Delivery App', 'Campus Event App'],
    'AI Engineer': ['AI Resume Analyzer', 'AI Chat Assistant', 'Recommendation System', 'Intelligent Document Processor'],
    'NLP Engineer': ['Sentiment Analysis System', 'Resume NLP Analyzer', 'Question Answering System', 'Text Classification API'],
    'Computer Vision Engineer': ['Number Plate Recognition', 'Object Detection App', 'Image Classification System', 'Document Image Analyzer'],
    'Cloud Architect': ['Multi-tier Cloud Architecture', 'Highly Available E-commerce System', 'Serverless Data Platform', 'Disaster Recovery Architecture'],
    'Security Analyst': ['SOC Monitoring Dashboard', 'Security Event Analyzer', 'Incident Response Tracker', 'Threat Intelligence Dashboard'],
    'Data Engineer': ['ETL Data Pipeline', 'Real-time Kafka Analytics', 'Airflow Data Warehouse Pipeline', 'Spark Customer Analytics'],
    'MLOps Engineer': ['End-to-End ML Pipeline', 'Model Serving Platform', 'Kubernetes ML Deployment', 'ML Monitoring Dashboard'],
    'Site Reliability Engineer': ['Service Health Dashboard', 'Kubernetes Reliability Lab', 'Auto-healing Web Service', 'Observability Platform'],
    'QA Engineer': ['E-commerce Test Plan', 'API Testing Suite', 'Web Regression Test Suite', 'Bug Tracking Dashboard'],
    'Automation Test Engineer': ['Selenium E-commerce Automation', 'Pytest API Automation', 'Web UI Regression Automation', 'CI/CD Test Pipeline'],
    'Network Engineer': ['Campus Network Design', 'Network Monitoring Dashboard', 'VPN Lab', 'Automated Network Configuration Tool'],
    'Systems Engineer': ['Server Monitoring Tool', 'Linux Automation Suite', 'Virtual Lab Environment', 'System Health Dashboard'],
    'Database Administrator': ['Database Backup System', 'Performance Monitoring Dashboard', 'High Availability Database Lab', 'Database Security Audit Tool'],
    'Solutions Architect': ['Scalable E-commerce Architecture', 'Cloud Migration Plan', 'Microservices Platform', 'Enterprise API Architecture'],
    'Generative AI Engineer': ['RAG Knowledge Assistant', 'Resume Copilot', 'Document Q&A System', 'Enterprise LLM Assistant'],
    'AI Application Engineer': ['AI Customer Support App', 'Resume Career Copilot', 'AI Document Assistant', 'LLM-powered Knowledge Portal'],
    'Robotics Engineer': ['Line Following Robot', 'ROS Navigation Robot', 'Vision-based Pick and Place', 'Autonomous Mobile Robot'],
    'Embedded Systems Engineer': ['Smart Sensor Controller', 'RTOS-based Device', 'Embedded Health Monitor', 'Microcontroller Automation System'],
    'Firmware Engineer': ['IoT Device Firmware', 'Motor Controller Firmware', 'Sensor Driver Library', 'Bootloader Prototype'],
    'IoT Engineer': ['Smart Home System', 'IoT Environmental Monitor', 'Industrial Sensor Dashboard', 'Connected Agriculture System'],
    'Hardware Engineer': ['Temperature Monitoring Board', 'Digital Clock PCB', 'IoT Sensor Board', 'Embedded Controller Prototype'],
    'Blockchain Developer': ['Token Management DApp', 'Decentralized Voting App', 'NFT Marketplace Prototype', 'Blockchain Certificate Verification'],
    'Game Developer': ['2D Platformer', '3D Adventure Game', 'Puzzle Game', 'Multiplayer Game Prototype'],
    'AR/VR Developer': ['AR Product Visualizer', 'VR Campus Tour', 'AR Learning App', 'Virtual Training Simulation'],
    'Platform Engineer': ['Internal Developer Portal', 'Kubernetes Platform', 'Self-service Deployment System', 'Developer Observability Platform'],
}


# =========================================================
# ML MODEL
# =========================================================

model = None
vectorizer = None

try:

    if MODEL_PATH.exists() and VECTORIZER_PATH.exists():

        model = joblib.load(MODEL_PATH)

        vectorizer = joblib.load(
            VECTORIZER_PATH
        )

        logger.info(
            "Auxiliary career ML model loaded."
        )

except Exception as exc:

    logger.warning(
        "Auxiliary ML model could not be loaded: %s",
        exc
    )


# =========================================================
# DATABASE
# =========================================================

def db_connection():

    return mysql.connector.connect(

        host=os.getenv(
            "DB_HOST",
            "localhost"
        ),

        port=int(
            os.getenv(
                "DB_PORT",
                "3306"
            )
        ),

        user=os.getenv(
            "DB_USER",
            "root"
        ),

        password=os.getenv(
            "DB_PASSWORD",
            ""
        ),

        database=os.getenv(
            "DB_NAME",
            "ai_resume_analyzer"
        ),

        connection_timeout=10,
    )


# =========================================================
# LOGIN DECORATOR
# =========================================================

def login_required(view):

    @wraps(view)
    def wrapped(*args, **kwargs):

        if "user_id" not in session:

            flash(
                "Please login to continue.",
                "error"
            )

            return redirect(
                url_for("login")
            )

        return view(
            *args,
            **kwargs
        )

    return wrapped


# =========================================================
# EMAIL HELPERS
# =========================================================

def valid_email(email):

    return bool(
        EMAIL_RE.fullmatch(
            email or ""
        )
    )


def valid_gmail(email):

    return bool(
        GMAIL_RE.fullmatch(
            email or ""
        )
    )


# =========================================================
# TEXT HELPERS
# =========================================================

def normalize(text):

    text = str(
        text or ""
    ).lower()

    text = text.replace(
        "\u00a0",
        " "
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


def clean_text(text):

    text = str(
        text or ""
    )

    text = text.replace(
        "\x00",
        " "
    )

    text = text.replace(
        "\r",
        "\n"
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n\s*\n\s*\n+",
        "\n\n",
        text
    )

    return text.strip()


def unique(items):

    seen = set()
    output = []

    for item in items:

        item = str(
            item
        ).strip()

        key = item.lower()

        if item and key not in seen:

            seen.add(key)

            output.append(
                item
            )

    return output


# =========================================================
# FILE HELPERS
# =========================================================

def allowed_file(filename):

    return (
        Path(filename)
        .suffix
        .lower()
        in ALLOWED_EXTENSIONS
    )


# =========================================================
# PDF EXTRACTION
# =========================================================

def extract_pdf(path):

    parts = []

    with pymupdf.open(path) as pdf:

        for page in pdf:

            txt = page.get_text(
                "text"
            )

            if txt:

                parts.append(
                    txt
                )

    text = clean_text(
        "\n".join(parts)
    )

    if len(text) >= 50:

        return text

    # OCR fallback
    ocr_parts = []

    with pymupdf.open(path) as pdf:

        for page in pdf:

            pix = page.get_pixmap(
                matrix=pymupdf.Matrix(
                    2,
                    2
                ),
                alpha=False
            )

            image = Image.open(
                io.BytesIO(
                    pix.tobytes("png")
                )
            )

            try:

                ocr_parts.append(
                    pytesseract.image_to_string(
                        image,
                        config="--oem 3 --psm 6"
                    )
                )

            finally:

                image.close()

    return clean_text(
        "\n".join(
            ocr_parts
        )
    )


# =========================================================
# DOCX EXTRACTION
# =========================================================

def extract_docx(path):

    doc = Document(
        path
    )

    parts = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip()
    ]

    for table in doc.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:

                parts.append(
                    " | ".join(cells)
                )

    return clean_text(
        "\n".join(parts)
    )


# =========================================================
# RESUME EXTRACTION
# =========================================================

def extract_resume(path):

    suffix = (
        Path(path)
        .suffix
        .lower()
    )

    if suffix == ".pdf":

        return extract_pdf(
            path
        )

    if suffix == ".docx":

        return extract_docx(
            path
        )

    return ""


# =========================================================
# SKILL DETECTION
# =========================================================

def detect_skills(text):

    normalized = normalize(
        text
    )

    found = []

    for skill in SKILLS:

        aliases = [
            skill
        ]

        aliases += [
            alias
            for alias, canonical
            in SKILL_ALIASES.items()
            if canonical == skill
        ]

        for alias in aliases:

            pattern = (
                r"(?<![a-z0-9])"
                +
                re.escape(
                    normalize(alias)
                )
                +
                r"(?![a-z0-9])"
            )

            if re.search(
                pattern,
                normalized
            ):

                found.append(
                    skill
                )

                break

    return unique(
        found
    )


# =========================================================
# SECTION EXTRACTION
# =========================================================

def section_lines(text, section):

    all_aliases = {

        "education": [
            "education",
            "educational background",
            "academic background",
            "qualifications"
        ],

        "projects": [
            "projects",
            "academic projects",
            "personal projects",
            "project experience"
        ],

        "experience": [
            "experience",
            "work experience",
            "professional experience",
            "employment",
            "internship",
            "internships"
        ],

        "certifications": [
            "certifications",
            "certification",
            "certificates",
            "courses",
            "training"
        ],

        "summary": [
            "summary",
            "professional summary",
            "profile",
            "objective",
            "career objective"
        ],
    }

    wanted = {
        normalize(x)
        for x in all_aliases.get(
            section,
            []
        )
    }

    headings = {
        normalize(x)
        for values in all_aliases.values()
        for x in values
    }

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    output = []
    started = False

    for line in lines:

        heading = normalize(
            re.sub(
                r"^[\dIVXivx]+[.)\-:]\s*",
                "",
                line
            ).rstrip(
                ":|-–— "
            )
        )

        if heading in wanted:

            started = True
            continue

        if started and heading in headings:

            break

        if started and len(line) > 1:

            output.append(
                re.sub(
                    r"^[•●▪◦■□◆◇*\-]+\s*",
                    "",
                    line
                ).strip()
            )

    return unique(
        output
    )[:40]


# =========================================================
# EDUCATION
# =========================================================

def extract_education(text):

    results = section_lines(
        text,
        "education"
    )

    if results:

        return results

    patterns = (
        r"\b("
        r"b\.?tech|"
        r"m\.?tech|"
        r"bca|"
        r"mca|"
        r"b\.?sc|"
        r"m\.?sc|"
        r"bachelor|"
        r"master|"
        r"engineering|"
        r"university|"
        r"college|"
        r"diploma"
        r")\b"
    )

    return unique(
        [
            line.strip()
            for line in text.splitlines()
            if re.search(
                patterns,
                normalize(line)
            )
        ]
    )[:30]


# =========================================================
# OTHER EXTRACTIONS
# =========================================================

def extract_projects(text):

    return section_lines(
        text,
        "projects"
    )


def extract_experience(text):

    return section_lines(
        text,
        "experience"
    )


def extract_certifications(text):

    return section_lines(
        text,
        "certifications"
    )


# =========================================================
# RESUME SCORE
# =========================================================

def resume_score(
    skills,
    education,
    projects,
    experience,
    certifications
):

    score = (
        min(
            len(skills) * 3,
            30
        )
        +
        (20 if education else 0)
        +
        min(
            len(projects) * 5,
            20
        )
        +
        (20 if experience else 0)
        +
        (10 if certifications else 0)
    )

    return min(
        score,
        100
    )


# =========================================================
# CAREER MATCH
# =========================================================

def career_match(
    role,
    skills
):

    required = CAREERS[
        role
    ]

    found = {
        normalize(s)
        for s in skills
    }

    matched = [
        s
        for s in required
        if normalize(s) in found
    ]

    missing = [
        s
        for s in required
        if normalize(s) not in found
    ]

    score = (
        round(
            (
                len(matched)
                /
                len(required)
            )
            * 100,
            2
        )
        if required
        else 0
    )

    return (
        score,
        matched,
        missing
    )


# =========================================================
# ML PREDICTION
# =========================================================

def ml_prediction(text):

    if (
        model is None
        or vectorizer is None
    ):

        return (
            None,
            0.0
        )

    try:

        x = vectorizer.transform(
            [text]
        )

        prediction = str(
            model.predict(x)[0]
        )

        if hasattr(
            model,
            "predict_proba"
        ):

            confidence = (
                float(
                    max(
                        model.predict_proba(x)[0]
                    )
                )
                * 100
            )

        else:

            confidence = 0.0

        return (
            prediction,
            round(
                confidence,
                2
            )
        )

    except Exception:

        return (
            None,
            0.0
        )


# =========================================================
# OTP EMAIL
# =========================================================

def send_otp(
    email,
    otp,
    purpose
):

    if (
        not app.config.get(
            "MAIL_USERNAME"
        )
        or
        not app.config.get(
            "MAIL_PASSWORD"
        )
    ):

        logger.error(
            "MAIL_USERNAME/MAIL_PASSWORD missing"
        )

        return False

    if purpose == "register":

        subject = (
            "AI Resume Analyzer - "
            "Email Verification OTP"
        )

    else:

        subject = (
            "AI Resume Analyzer - "
            "Password Reset OTP"
        )

    body = f"""
Hello,

Your OTP is: {otp}

This OTP is valid for 90 seconds.

If you did not request this OTP,
please ignore this email.

Regards,
AI Resume Analyzer
"""

    try:

        message = Message(
            subject=subject,
            sender=app.config[
                "MAIL_USERNAME"
            ],
            recipients=[
                email
            ],
            body=body
        )

        mail.send(
            message
        )

        return True

    except Exception as exc:

        logger.exception(
            "OTP email failed: %s",
            exc
        )

        return False


# =========================================================
# ISSUE OTP
# =========================================================

def issue_otp(
    email,
    purpose
):

    otp = (
        f"{secrets.randbelow(1_000_000):06d}"
    )

    now = datetime.utcnow()

    expires = (
        now
        +
        timedelta(
            seconds=app.config[
                "OTP_EXPIRY_SECONDS"
            ]
        )
    )

    db = None
    cursor = None
    otp_id = None

    try:

        db = db_connection()

        cursor = db.cursor()

        # Invalidate old OTPs
        cursor.execute(
            """
            UPDATE otp_tokens
            SET consumed_at = UTC_TIMESTAMP()
            WHERE email = %s
              AND purpose = %s
              AND consumed_at IS NULL
            """,
            (
                email,
                purpose
            )
        )

        # Insert new OTP hash
        cursor.execute(
            """
            INSERT INTO otp_tokens
            (
                email,
                purpose,
                otp_hash,
                expires_at,
                attempts
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                0
            )
            """,
            (
                email,
                purpose,
                generate_password_hash(
                    otp
                ),
                expires
            )
        )

        otp_id = cursor.lastrowid

        db.commit()

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()

    # Send email
    if not send_otp(
        email,
        otp,
        purpose
    ):

        db = None
        cursor = None

        try:

            db = db_connection()

            cursor = db.cursor()

            if otp_id:

                cursor.execute(
                    """
                    DELETE FROM otp_tokens
                    WHERE id = %s
                    """,
                    (
                        otp_id,
                    )
                )

                db.commit()

        except Exception:

            logger.exception(
                "Could not clean failed OTP record"
            )

        finally:

            if cursor:

                cursor.close()

            if (
                db
                and
                db.is_connected()
            ):

                db.close()

        return (
            False,
            None
        )

    return (
        True,
        expires
    )


# =========================================================
# VERIFY OTP
# =========================================================

def verify_otp(
    email,
    purpose,
    otp
):

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT *
            FROM otp_tokens
            WHERE email = %s
              AND purpose = %s
              AND consumed_at IS NULL
            ORDER BY id DESC
            LIMIT 1
            """,
            (
                email,
                purpose
            )
        )

        row = cursor.fetchone()

        if not row:

            return (
                False,
                "OTP not found. Please request a new OTP."
            )

        if datetime.utcnow() > row[
            "expires_at"
        ]:

            return (
                False,
                "OTP expired. Please request a new OTP."
            )

        if row[
            "attempts"
        ] >= app.config[
            "OTP_MAX_ATTEMPTS"
        ]:

            return (
                False,
                "Too many incorrect attempts. Please request a new OTP."
            )

        if not check_password_hash(
            row["otp_hash"],
            otp
        ):

            cursor.execute(
                """
                UPDATE otp_tokens
                SET attempts = attempts + 1
                WHERE id = %s
                """,
                (
                    row["id"],
                )
            )

            db.commit()

            return (
                False,
                "Invalid OTP. Please try again."
            )

        cursor.execute(
            """
            UPDATE otp_tokens
            SET consumed_at = UTC_TIMESTAMP()
            WHERE id = %s
            """,
            (
                row["id"],
            )
        )

        db.commit()

        return (
            True,
            "OTP verified successfully."
        )

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()


# =========================================================
# LATEST OTP
# =========================================================

def latest_otp_created(
    email,
    purpose
):

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT created_at
            FROM otp_tokens
            WHERE email = %s
              AND purpose = %s
            ORDER BY id DESC
            LIMIT 1
            """,
            (
                email,
                purpose
            )
        )

        row = cursor.fetchone()

        return (
            row["created_at"]
            if row
            else None
        )

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()


# =========================================================
# SAVE RESUME + ANALYSIS
# =========================================================

def save_resume_and_analysis(
    filename,
    stored_filename,
    role,
    score,
    match,
    skills,
    missing,
    resume_score_breakdown
):

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor()

        cursor.execute(
            """
            INSERT INTO resumes
            (
                user_id,
                original_filename,
                stored_filename
            )
            VALUES
            (
                %s,
                %s,
                %s
            )
            """,
            (
                session["user_id"],
                filename,
                stored_filename
            )
        )

        resume_id = cursor.lastrowid

        cursor.execute(
            """
            INSERT INTO resume_analysis
            (
                user_id,
                resume_id,
                filename,
                target_career,
                resume_score,
                match_percentage,
                matched_skills,
                missing_skills,
                score_breakdown
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )
            """,
            (
                session["user_id"],
                resume_id,
                filename,
                role,
                score,
                match,
                ", ".join(skills),
                ", ".join(missing),
                str(
                    resume_score_breakdown
                )
            )
        )

        analysis_id = cursor.lastrowid

        db.commit()

        return (
            resume_id,
            analysis_id
        )

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()


# =========================================================
# SAVE EXISTING ANALYSIS
# =========================================================

def save_existing_analysis(
    resume_id,
    filename,
    role,
    score,
    match,
    skills,
    missing,
    breakdown
):

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT id
            FROM resumes
            WHERE id = %s
              AND user_id = %s
            """,
            (
                resume_id,
                session["user_id"]
            )
        )

        if not cursor.fetchone():

            return None

        cursor.execute(
            """
            INSERT INTO resume_analysis
            (
                user_id,
                resume_id,
                filename,
                target_career,
                resume_score,
                match_percentage,
                matched_skills,
                missing_skills,
                score_breakdown
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )
            """,
            (
                session["user_id"],
                resume_id,
                filename,
                role,
                score,
                match,
                ", ".join(skills),
                ", ".join(missing),
                str(breakdown)
            )
        )

        analysis_id = cursor.lastrowid

        db.commit()

        return analysis_id

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()


# =========================================================
# ANALYZE TEXT
# =========================================================

def analyze_text(
    text,
    role
):

    skills = detect_skills(
        text
    )

    education = extract_education(
        text
    )

    projects = extract_projects(
        text
    )

    experience = extract_experience(
        text
    )

    certifications = extract_certifications(
        text
    )

    score = resume_score(
        skills,
        education,
        projects,
        experience,
        certifications
    )

    match, matched, missing = career_match(
        role,
        skills
    )

    predicted, confidence = ml_prediction(
        text
    )

    priority = []

    for index, skill in enumerate(
        missing
    ):

        if index < max(
            1,
            len(missing) // 2
        ):

            level = "High"

        elif index < max(
            2,
            len(missing) - 1
        ):

            level = "Medium"

        else:

            level = "Low"

        priority.append(
            {
                "skill": skill,
                "priority": level
            }
        )

    topics = []

    for skill in missing:

        topics.append(
            {
                "skill": skill,
                "topics": LEARNING_TOPICS.get(
                    skill,
                    [
                        f"{skill} fundamentals",
                        f"{skill} practical usage",
                        f"{skill} projects"
                    ]
                )
            }
        )

    return {

        "skills": skills,

        "education": education,

        "projects": projects,

        "experience": experience,

        "certifications": certifications,

        "resume_score": score,

        "match_percentage": match,

        "matched_skills": matched,

        "missing_skills": missing,

        "priority": priority,

        "predicted_career": predicted,

        "ml_confidence": confidence,

        "topics": topics,

        "roadmap": CAREER_ROADMAPS.get(
            role,
            CAREER_ROADMAPS[
                "Software Engineer"
            ]
        ),

        "projects_recommended": PROJECTS.get(
            role,
            [
                f"Build a beginner {role} project",
                f"Build an intermediate {role} project",
                f"Build an industry-style {role} project"
            ]
        ),
    }


# =========================================================
# HOME
# =========================================================

@app.route("/")
def home():

    if "user_id" not in session:

        return redirect(
            url_for("login")
        )

    return render_template(
        "index.html",
        careers=sorted(
            CAREERS
        ),
        recent_resumes=get_user_resumes()
    )


# =========================================================
# LOGIN
# =========================================================

@app.route(
    "/login",
    methods=[
        "GET",
        "POST"
    ]
)
def login():

    if request.method == "POST":

        email = (
            request.form
            .get(
                "email",
                ""
            )
            .strip()
            .lower()
        )

        password = request.form.get(
            "password",
            ""
        )

        if (
            not valid_email(email)
            or
            not password
        ):

            flash(
                "Enter a valid email and password.",
                "error"
            )

            return render_template(
                "login.html"
            )

        db = None
        cursor = None

        try:

            db = db_connection()

            cursor = db.cursor(
                dictionary=True
            )

            cursor.execute(
                """
                SELECT id, name, email, password
                FROM users
                WHERE email = %s
                LIMIT 1
                """,
                (
                    email,
                )
            )

            user = cursor.fetchone()

            if (
                user
                and
                check_password_hash(
                    user["password"],
                    password
                )
            ):

                session.clear()

                session["user_id"] = user[
                    "id"
                ]

                session["user_name"] = user[
                    "name"
                ]

                session["user_email"] = user[
                    "email"
                ]

                return redirect(
                    url_for("home")
                )

            flash(
                "Invalid email or password.",
                "error"
            )

        except Error:

            logger.exception(
                "Login error"
            )

            flash(
                "Unable to login right now. Check database configuration.",
                "error"
            )

        finally:

            if cursor:

                cursor.close()

            if (
                db
                and
                db.is_connected()
            ):

                db.close()

    return render_template(
        "login.html"
    )


# =========================================================
# REGISTER
# =========================================================

@app.route(
    "/register",
    methods=[
        "GET",
        "POST"
    ]
)
def register():

    # -----------------------------------------------------
    # GET REGISTER PAGE
    # -----------------------------------------------------

    if request.method == "GET":

        pending = session.get(
            "pending_registration"
        )

        expires_at = session.get(
            "otp_expires_at",
            0
        )

        now = datetime.utcnow().timestamp()

        otp_sent = bool(
            pending
            and
            expires_at > now
            and
            not session.get(
                "registration_verified"
            )
        )

        remaining = 0

        if otp_sent:

            remaining = max(
                0,
                int(
                    expires_at - now
                )
            )

        return render_template(
            "register.html",

            pending_name=(
                pending.get(
                    "name",
                    ""
                )
                if pending
                else ""
            ),

            pending_email=(
                pending.get(
                    "email",
                    ""
                )
                if pending
                else ""
            ),

            otp_sent=otp_sent,

            otp_remaining=remaining
        )

    # -----------------------------------------------------
    # POST - SEND OTP
    # -----------------------------------------------------

    name = (
        request.form
        .get(
            "name",
            ""
        )
        .strip()
    )

    email = (
        request.form
        .get(
            "email",
            ""
        )
        .strip()
        .lower()
    )

    if len(name) < 2:

        flash(
            "Enter your full name.",
            "error"
        )

        return render_template(
            "register.html",
            pending_name=name,
            pending_email=email,
            otp_sent=False,
            otp_remaining=0
        )

    # Gmail only
    if not valid_gmail(email):

        flash(
            "Please enter a valid Gmail address.",
            "error"
        )

        return render_template(
            "register.html",
            pending_name=name,
            pending_email=email,
            otp_sent=False,
            otp_remaining=0
        )

    # -----------------------------------------------------
    # CHECK DUPLICATE EMAIL
    # -----------------------------------------------------

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT id
            FROM users
            WHERE email = %s
            LIMIT 1
            """,
            (
                email,
            )
        )

        existing_user = cursor.fetchone()

        if existing_user:

            flash(
                "Email already registered. Please login.",
                "error"
            )

            return render_template(
                "login.html",
                prefill_email=email
            )

    except Error:

        logger.exception(
            "Registration email check failed"
        )

        flash(
            "Unable to check email right now.",
            "error"
        )

        return render_template(
            "register.html",
            pending_name=name,
            pending_email=email,
            otp_sent=False,
            otp_remaining=0
        )

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()

    # -----------------------------------------------------
    # STORE PENDING REGISTRATION
    # -----------------------------------------------------

    session["pending_registration"] = {
        "name": name,
        "email": email
    }

    session.pop(
        "registration_verified",
        None
    )

    # -----------------------------------------------------
    # SEND OTP
    # -----------------------------------------------------

    ok, expires = issue_otp(
        email,
        "register"
    )

    if not ok:

        flash(
            "OTP could not be sent. Please try again.",
            "error"
        )

        return render_template(
            "register.html",
            pending_name=name,
            pending_email=email,
            otp_sent=False,
            otp_remaining=0
        )

    # -----------------------------------------------------
    # SERVER-SIDE EXPIRY
    # -----------------------------------------------------

    session["otp_expires_at"] = (
        expires.timestamp()
    )

    remaining = max(
        0,
        int(
            expires.timestamp()
            -
            datetime.utcnow().timestamp()
        )
    )

    # IMPORTANT:
    # Same register page
    return render_template(
        "register.html",

        pending_name=name,

        pending_email=email,

        otp_sent=True,

        otp_remaining=remaining
    )


# =========================================================
# REGISTER OTP VERIFY
# =========================================================

@app.route(
    "/register/verify-otp",
    methods=["POST"]
)
def register_verify_otp():

    pending = session.get(
        "pending_registration"
    )

    if not pending:

        return jsonify(
            success=False,
            message=(
                "Registration session expired. "
                "Please start again."
            )
        ), 400

    email = pending[
        "email"
    ]

    otp = (
        request.form
        .get(
            "otp",
            ""
        )
        .strip()
    )

    if not re.fullmatch(
        r"\d{6}",
        otp
    ):

        return jsonify(
            success=False,
            message="Enter the 6-digit OTP."
        ), 400

    # Server-side expiry
    expires_at = session.get(
        "otp_expires_at",
        0
    )

    if (
        datetime.utcnow().timestamp()
        >
        expires_at
    ):

        return jsonify(
            success=False,
            message=(
                "OTP expired. "
                "Please request a new OTP."
            )
        ), 400

    success, message = verify_otp(
        email,
        "register",
        otp
    )

    if not success:

        return jsonify(
            success=False,
            message=message
        ), 400

    # OTP verified
    session["registration_verified"] = True

    # OTP should no longer be considered active
    session["otp_expires_at"] = 0

    return jsonify(
        success=True,
        message="OTP verified successfully."
    )


# =========================================================
# RESEND OTP
# =========================================================

@app.route(
    "/resend-otp",
    methods=["POST"]
)
def resend_otp():

    purpose = request.form.get(
        "purpose",
        "register"
    )

    if purpose not in {
        "register",
        "reset"
    }:

        return jsonify(
            success=False,
            message="Invalid OTP request."
        ), 400

    if purpose == "register":

        pending = session.get(
            "pending_registration"
        )

    else:

        pending = session.get(
            "pending_reset"
        )

    if not pending:

        return jsonify(
            success=False,
            message="Verification session expired."
        ), 400

    email = pending[
        "email"
    ]

    # Cooldown
    last = latest_otp_created(
        email,
        purpose
    )

    if last:

        elapsed = (
            datetime.utcnow()
            -
            last
        ).total_seconds()

        if elapsed < app.config[
            "OTP_RESEND_COOLDOWN_SECONDS"
        ]:

            wait = (
                app.config[
                    "OTP_RESEND_COOLDOWN_SECONDS"
                ]
                -
                int(elapsed)
            )

            return jsonify(
                success=False,
                message=(
                    f"Please wait "
                    f"{max(wait, 1)} "
                    f"seconds before requesting another OTP."
                )
            ), 429

    # Issue new OTP
    ok, expires = issue_otp(
        email,
        purpose
    )

    if not ok:

        return jsonify(
            success=False,
            message="Unable to send OTP."
        ), 500

    session["otp_expires_at"] = (
        expires.timestamp()
    )

    if purpose == "register":

        session.pop(
            "registration_verified",
            None
        )

    return jsonify(
        success=True,
        message="New OTP sent.",
        expires_at=expires.timestamp()
    )


# =========================================================
# CREATE PASSWORD
# =========================================================

@app.route(
    "/create-password",
    methods=[
        "GET",
        "POST"
    ]
)
def create_password():

    pending = session.get(
        "pending_registration"
    )

    if (
        not pending
        or
        not session.get(
            "registration_verified"
        )
    ):

        return redirect(
            url_for("register")
        )

    if request.method == "POST":

        password = request.form.get(
            "password",
            ""
        )

        confirm = request.form.get(
            "confirm_password",
            ""
        )

        if (
            len(password) < 8
            or
            password != confirm
        ):

            flash(
                "Password must be at least 8 characters and both passwords must match.",
                "error"
            )

            return render_template(
                "create_password.html"
            )

        db = None
        cursor = None

        try:

            db = db_connection()

            cursor = db.cursor()

            cursor.execute(
                """
                INSERT INTO users
                (
                    name,
                    email,
                    password
                )
                VALUES
                (
                    %s,
                    %s,
                    %s
                )
                """,
                (
                    pending["name"],
                    pending["email"],
                    generate_password_hash(
                        password
                    )
                )
            )

            db.commit()

            session.clear()

            flash(
                "Account created successfully. Please login.",
                "success"
            )

            return redirect(
                url_for("login")
            )

        except mysql.connector.IntegrityError:

            if db:

                db.rollback()

            flash(
                "Email already registered. Please login.",
                "error"
            )

            return redirect(
                url_for("login")
            )

        except Error:

            if db:

                db.rollback()

            flash(
                "Unable to create account right now.",
                "error"
            )

        finally:

            if cursor:

                cursor.close()

            if (
                db
                and
                db.is_connected()
            ):

                db.close()

    return render_template(
        "create_password.html"
    )


# =========================================================
# FORGOT PASSWORD
# =========================================================

@app.route(
    "/forgot-password",
    methods=[
        "GET",
        "POST"
    ]
)
def forgot_password():

    if request.method == "POST":

        email = (
            request.form
            .get(
                "email",
                ""
            )
            .strip()
            .lower()
        )

        if not valid_email(email):

            flash(
                "Enter a valid email.",
                "error"
            )

            return render_template(
                "forgot_password.html"
            )

        db = None
        cursor = None
        user = None

        try:

            db = db_connection()

            cursor = db.cursor(
                dictionary=True
            )

            cursor.execute(
                """
                SELECT id
                FROM users
                WHERE email = %s
                LIMIT 1
                """,
                (
                    email,
                )
            )

            user = cursor.fetchone()

        finally:

            if cursor:

                cursor.close()

            if (
                db
                and
                db.is_connected()
            ):

                db.close()

        if user:

            last = latest_otp_created(
                email,
                "reset"
            )

            if (
                last
                and
                (
                    datetime.utcnow()
                    -
                    last
                ).total_seconds()
                <
                app.config[
                    "OTP_RESEND_COOLDOWN_SECONDS"
                ]
            ):

                flash(
                    "If the account exists, please wait before requesting another code.",
                    "success"
                )

                return render_template(
                    "forgot_password.html"
                )

            ok, expires = issue_otp(
                email,
                "reset"
            )

            if ok:

                session["pending_reset"] = {
                    "email": email
                }

                session["reset_verified"] = False

                session["otp_expires_at"] = (
                    expires.timestamp()
                )

                return redirect(
                    url_for(
                        "verify_otp_page",
                        purpose="reset"
                    )
                )

        flash(
            "If an account exists for that email, a password-reset OTP has been sent.",
            "success"
        )

    return render_template(
        "forgot_password.html"
    )


# =========================================================
# RESET OTP PAGE
# =========================================================

@app.route(
    "/verify-otp/<purpose>",
    methods=["GET"]
)
def verify_otp_page(
    purpose
):

    if purpose != "reset":

        return redirect(
            url_for("login")
        )

    pending = session.get(
        "pending_reset"
    )

    if not pending:

        return redirect(
            url_for("forgot_password")
        )

    expires_at = session.get(
        "otp_expires_at",
        0
    )

    remaining = max(
        0,
        int(
            expires_at
            -
            datetime.utcnow().timestamp()
        )
    )

    return render_template(
        "verify_otp.html",
        purpose=purpose,
        email=pending["email"],
        otp_remaining=remaining
    )


# =========================================================
# RESET OTP VERIFY
# =========================================================

@app.route(
    "/verify-reset-otp",
    methods=["POST"]
)
def verify_reset_otp():

    pending = session.get(
        "pending_reset"
    )

    if not pending:

        return jsonify(
            success=False,
            message="Verification session expired."
        ), 400

    otp = (
        request.form
        .get(
            "otp",
            ""
        )
        .strip()
    )

    if not re.fullmatch(
        r"\d{6}",
        otp
    ):

        return jsonify(
            success=False,
            message="Enter the 6-digit OTP."
        ), 400

    expires_at = session.get(
        "otp_expires_at",
        0
    )

    if (
        datetime.utcnow().timestamp()
        >
        expires_at
    ):

        return jsonify(
            success=False,
            message="OTP expired. Please request a new OTP."
        ), 400

    success, message = verify_otp(
        pending["email"],
        "reset",
        otp
    )

    if not success:

        return jsonify(
            success=False,
            message=message
        ), 400

    session["reset_verified"] = True

    session["otp_expires_at"] = 0

    return jsonify(
        success=True,
        message="OTP verified successfully."
    )


# =========================================================
# RESET PASSWORD
# =========================================================

@app.route(
    "/reset-password",
    methods=[
        "GET",
        "POST"
    ]
)
def reset_password():

    pending = session.get(
        "pending_reset"
    )

    if (
        not pending
        or
        not session.get(
            "reset_verified"
        )
    ):

        return redirect(
            url_for(
                "forgot_password"
            )
        )

    if request.method == "POST":

        password = request.form.get(
            "password",
            ""
        )

        confirm = request.form.get(
            "confirm_password",
            ""
        )

        if (
            len(password) < 8
            or
            password != confirm
        ):

            flash(
                "Password must be at least 8 characters and both passwords must match.",
                "error"
            )

            return render_template(
                "reset_password.html"
            )

        db = None
        cursor = None

        try:

            db = db_connection()

            cursor = db.cursor()

            cursor.execute(
                """
                UPDATE users
                SET password = %s
                WHERE email = %s
                """,
                (
                    generate_password_hash(
                        password
                    ),
                    pending["email"]
                )
            )

            db.commit()

            session.clear()

            flash(
                "Password changed successfully. Please login.",
                "success"
            )

            return redirect(
                url_for("login")
            )

        except Error:

            if db:

                db.rollback()

            flash(
                "Unable to reset password right now.",
                "error"
            )

        finally:

            if cursor:

                cursor.close()

            if (
                db
                and
                db.is_connected()
            ):

                db.close()

    return render_template(
        "reset_password.html"
    )


# =========================================================
# USER RESUMES
# =========================================================

def get_user_resumes():

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT
                id,
                original_filename,
                created_at
            FROM resumes
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (
                session["user_id"],
            )
        )

        return cursor.fetchall()

    except Error:

        return []

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()


# =========================================================
# RESULT CONTEXT
# =========================================================

def build_result_context(
    analysis,
    role
):

    learning_languages = [

        skill

        for skill in CAREERS[
            role
        ]

        if skill in {
            "python",
            "java",
            "javascript",
            "typescript",
            "c++",
            "c",
            "kotlin",
            "dart"
        }
    ]

    return {
        **analysis,
        "role": role,
        "learning_languages": learning_languages
    }


# =========================================================
# ANALYZE NEW RESUME
# =========================================================

@app.route(
    "/analyze",
    methods=["POST"]
)
@login_required
def analyze():

    role = (
        request.form
        .get(
            "target_role",
            ""
        )
        .strip()
    )

    if role not in CAREERS:

        flash(
            "Please choose a valid career.",
            "error"
        )

        return redirect(
            url_for("home")
        )

    file = request.files.get(
        "resume"
    )

    if (
        not file
        or
        not file.filename
    ):

        flash(
            "Please upload a PDF or DOCX resume.",
            "error"
        )

        return redirect(
            url_for("home")
        )

    # -------------------------------------------------
    # FILE TYPE CHECK
    # -------------------------------------------------

    original = secure_filename(
        file.filename
    )

    if not allowed_file(
        original
    ):

        flash(
            "Only PDF and DOCX files are supported.",
            "error"
        )

        return redirect(
            url_for("home")
        )

    stored = (
        f"{uuid.uuid4().hex}"
        f"{Path(original).suffix.lower()}"
    )

    path = (
        UPLOAD_FOLDER
        /
        stored
    )

    try:

        # -------------------------------------------------
        # SAVE FILE
        # -------------------------------------------------

        file.save(
            path
        )

        # -------------------------------------------------
        # EXTRACT CONTENT
        # -------------------------------------------------

        text = extract_resume(
            path
        )

        # -------------------------------------------------
        # BASIC TEXT CHECK
        # -------------------------------------------------

        if not text or len(
            text.strip()
        ) < 30:

            path.unlink(
                missing_ok=True
            )

            flash(
                "Invalid file. Please upload a readable resume.",
                "error"
            )

            return redirect(
                url_for("home")
            )

        # -------------------------------------------------
        # RESUME CONTENT VALIDATION
        # IMPORTANT:
        # File name is NOT used here.
        # We check the actual PDF/DOCX content.
        # -------------------------------------------------

        normalized_text = " ".join(
            text.lower().split()
        )

        # Common resume section keywords
        resume_sections = [
            "education",
            "skills",
            "technical skills",
            "projects",
            "experience",
            "work experience",
            "professional experience",
            "certifications",
            "professional summary",
            "summary",
            "objective",
            "career objective",
            "internship",
            "internships",
            "achievements",
            "extracurricular",
            "languages",
            "profile",
            "contact",
        ]

        # Count resume sections found inside the document
        matched_sections = sum(
            1
            for section in resume_sections
            if section in normalized_text
        )

        # -------------------------------------------------
        # CHECK FOR CONTACT INFORMATION
        # -------------------------------------------------

        has_email = bool(
            re.search(
                r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
                text
            )
        )

        has_phone = bool(
            re.search(
                r"(?:\+91[\s-]?)?[6-9]\d{9}\b",
                text
            )
        )

        # -------------------------------------------------
        # CHECK COMMON RESUME SKILLS
        # -------------------------------------------------

        skill_matches = 0

        for skill in SKILLS:

            skill_pattern = (
                r"\b"
                + re.escape(
                    skill.lower()
                )
                + r"\b"
            )

            if re.search(
                skill_pattern,
                normalized_text
            ):

                skill_matches += 1

        # -------------------------------------------------
        # RESUME VALIDATION LOGIC
        #
        # Accept when the document has:
        #
        # 1. At least 3 resume sections
        #
        # OR
        #
        # 2. At least 2 sections +
        #    contact information +
        #    at least 2 skills
        #
        # This means:
        # abc.pdf + resume content  -> ACCEPT
        # resume.pdf + random PDF   -> REJECT
        # -------------------------------------------------

        valid_resume = False

        if matched_sections >= 3:

            valid_resume = True

        elif (
            matched_sections >= 2
            and
            (has_email or has_phone)
            and
            skill_matches >= 2
        ):

            valid_resume = True

        # -------------------------------------------------
        # REJECT NON-RESUME DOCUMENT
        # -------------------------------------------------

        if not valid_resume:

            path.unlink(
                missing_ok=True
            )

            flash(
                "Invalid file. Please upload a valid resume.",
                "error"
            )

            return redirect(
                url_for("home")
            )

        # -------------------------------------------------
        # ANALYZE RESUME
        # -------------------------------------------------

        analysis = analyze_text(
            text,
            role
        )

        # -------------------------------------------------
        # RESUME SCORE BREAKDOWN
        # -------------------------------------------------

        breakdown = {

            "skills": min(
                len(
                    analysis["skills"]
                ) * 3,
                30
            ),

            "education": (
                20
                if analysis["education"]
                else 0
            ),

            "projects": min(
                len(
                    analysis["projects"]
                ) * 5,
                20
            ),

            "experience": (
                20
                if analysis["experience"]
                else 0
            ),

            "certifications": (
                10
                if analysis["certifications"]
                else 0
            ),
        }

        # -------------------------------------------------
        # SAVE ANALYSIS
        # -------------------------------------------------

        resume_id, analysis_id = (
            save_resume_and_analysis(
                original,
                stored,
                role,
                analysis["resume_score"],
                analysis["match_percentage"],
                analysis["matched_skills"],
                analysis["missing_skills"],
                breakdown
            )
        )

        # -------------------------------------------------
        # SHOW RESULT
        # -------------------------------------------------

        return render_template(
            "result.html",
            analysis_id=analysis_id,
            resume_id=resume_id,
            filename=original,
            result=build_result_context(
                analysis,
                role
            )
        )

    except Exception:

        logger.exception(
            "Analysis failed"
        )

        path.unlink(
            missing_ok=True
        )

        flash(
            "Unable to analyze this resume right now.",
            "error"
        )

        return redirect(
            url_for("home")
        )

# =========================================================
# ANALYZE EXISTING RESUME
# =========================================================

@app.route(
    "/analyze-existing",
    methods=["POST"]
)
@login_required
def analyze_existing():

    role = (
        request.form
        .get(
            "target_role",
            ""
        )
        .strip()
    )

    resume_id = (
        request.form
        .get(
            "resume_id",
            ""
        )
        .strip()
    )

    if (
        role not in CAREERS
        or
        not resume_id.isdigit()
    ):

        flash(
            "Invalid analysis request.",
            "error"
        )

        return redirect(
            url_for("home")
        )

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT
                id,
                original_filename,
                stored_filename
            FROM resumes
            WHERE id = %s
              AND user_id = %s
            """,
            (
                int(resume_id),
                session["user_id"]
            )
        )

        resume = cursor.fetchone()

        if not resume:

            flash(
                "Resume not found.",
                "error"
            )

            return redirect(
                url_for("home")
            )

        path = (
            UPLOAD_FOLDER
            /
            resume["stored_filename"]
        )

        if not path.exists():

            flash(
                "Stored resume file is unavailable. Please upload it again.",
                "error"
            )

            return redirect(
                url_for("home")
            )

        text = extract_resume(
            path
        )

        analysis = analyze_text(
            text,
            role
        )

        breakdown = {

            "skills": min(
                len(
                    analysis["skills"]
                ) * 3,
                30
            ),

            "education": (
                20
                if analysis["education"]
                else 0
            ),

            "projects": min(
                len(
                    analysis["projects"]
                ) * 5,
                20
            ),

            "experience": (
                20
                if analysis["experience"]
                else 0
            ),

            "certifications": (
                10
                if analysis["certifications"]
                else 0
            ),
        }

        analysis_id = save_existing_analysis(
            int(resume_id),
            resume["original_filename"],
            role,
            analysis["resume_score"],
            analysis["match_percentage"],
            analysis["matched_skills"],
            analysis["missing_skills"],
            breakdown
        )

        return render_template(
            "result.html",
            analysis_id=analysis_id,
            resume_id=int(resume_id),
            filename=resume[
                "original_filename"
            ],
            result=build_result_context(
                analysis,
                role
            )
        )

    except Exception:

        logger.exception(
            "Existing resume analysis failed"
        )

        flash(
            "Unable to analyze the selected resume.",
            "error"
        )

        return redirect(
            url_for("home")
        )

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()


# =========================================================
# HISTORY
# =========================================================

@app.route("/history")
@login_required
def history():

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT
                id,
                resume_id,
                filename,
                target_career,
                resume_score,
                match_percentage,
                created_at
            FROM resume_analysis
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (
                session["user_id"],
            )
        )

        analyses = cursor.fetchall()

        best = {}

        for row in analyses:

            rid = row[
                "resume_id"
            ]

            if (
                rid not in best
                or
                float(
                    row["match_percentage"]
                )
                >
                float(
                    best[rid][
                        "match_percentage"
                    ]
                )
            ):

                best[rid] = row

        for row in analyses:

            row["recommended"] = (
                best[
                    row["resume_id"]
                ]["id"]
                ==
                row["id"]
            )

        return render_template(
            "history.html",
            analyses=analyses
        )

    except Error:

        flash(
            "Unable to load history.",
            "error"
        )

        return render_template(
            "history.html",
            analyses=[]
        )

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()

# =========================================================
# CLEAR ALL HISTORY
# =========================================================

@app.route(
    "/clear-history",
    methods=["POST"]
)
@login_required
def clear_history():

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        user_id = session["user_id"]

        # -------------------------------------------------
        # Get all resumes belonging to this user
        # -------------------------------------------------

        cursor.execute(
            """
            SELECT id, stored_filename
            FROM resumes
            WHERE user_id = %s
            """,
            (
                user_id,
            )
        )

        resumes = cursor.fetchall()

        # -------------------------------------------------
        # Delete analysis history
        # -------------------------------------------------

        cursor.execute(
            """
            DELETE FROM resume_analysis
            WHERE user_id = %s
            """,
            (
                user_id,
            )
        )

        # -------------------------------------------------
        # Delete resume records
        # -------------------------------------------------

        cursor.execute(
            """
            DELETE FROM resumes
            WHERE user_id = %s
            """,
            (
                user_id,
            )
        )

        db.commit()

        # -------------------------------------------------
        # Delete uploaded resume files
        # -------------------------------------------------

        for resume in resumes:

            stored_filename = (
                resume.get(
                    "stored_filename"
                )
            )

            if stored_filename:

                file_path = (
                    UPLOAD_FOLDER
                    /
                    stored_filename
                )

                file_path.unlink(
                    missing_ok=True
                )

        flash(
            "All your history has been cleared successfully.",
            "success"
        )

        return redirect(
            url_for("history")
        )

    except Error:

        if db:
            db.rollback()

        logger.exception(
            "Clear history failed"
        )

        flash(
            "Unable to clear history.",
            "error"
        )

        return redirect(
            url_for("history")
        )

    finally:

        if cursor:
            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):
            db.close()

# =========================================================
# HISTORY DETAIL
# =========================================================

@app.route(
    "/history/<int:analysis_id>"
)
@login_required
def history_detail(
    analysis_id
):

    db = None
    cursor = None

    try:

        db = db_connection()

        cursor = db.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT
                a.*,
                r.stored_filename
            FROM resume_analysis a
            JOIN resumes r
              ON r.id = a.resume_id
             AND r.user_id = a.user_id
            WHERE a.id = %s
              AND a.user_id = %s
            """,
            (
                analysis_id,
                session["user_id"]
            )
        )

        row = cursor.fetchone()

        if not row:

            flash(
                "Analysis not found.",
                "error"
            )

            return redirect(
                url_for("history")
            )

        path = (
            UPLOAD_FOLDER
            /
            row["stored_filename"]
        )

        if not path.exists():

            flash(
                "The stored resume is unavailable. Please upload it again.",
                "error"
            )

            return redirect(
                url_for("history")
            )

        text = extract_resume(
            path
        )

        result = analyze_text(
            text,
            row["target_career"]
        )

        return render_template(
            "result.html",
            analysis_id=analysis_id,
            resume_id=row["resume_id"],
            filename=row["filename"],
            result=build_result_context(
                result,
                row["target_career"]
            ),
            from_history=True
        )

    except Exception:

        logger.exception(
            "History detail error"
        )

        flash(
            "Unable to open this analysis.",
            "error"
        )

        return redirect(
            url_for("history")
        )

    finally:

        if cursor:

            cursor.close()

        if (
            db
            and
            db.is_connected()
        ):

            db.close()


# =========================================================
# PROFILE
# =========================================================

@app.route(
    "/profile",
    methods=[
        "GET",
        "POST"
    ]
)
@login_required
def profile():

    if request.method == "POST":

        name = (
            request.form
            .get(
                "name",
                ""
            )
            .strip()
        )

        if len(name) < 2:

            flash(
                "Enter a valid name.",
                "error"
            )

            return redirect(
                url_for("profile")
            )

        db = None
        cursor = None

        try:

            db = db_connection()

            cursor = db.cursor()

            cursor.execute(
                """
                UPDATE users
                SET name = %s
                WHERE id = %s
                """,
                (
                    name,
                    session["user_id"]
                )
            )

            db.commit()

            session[
                "user_name"
            ] = name

            flash(
                "Profile updated.",
                "success"
            )

        finally:

            if cursor:

                cursor.close()

            if (
                db
                and
                db.is_connected()
            ):

                db.close()

    return render_template(
        "profile.html",
        name=session.get(
            "user_name"
        ),
        email=session.get(
            "user_email"
        )
    )


# =========================================================
# LOGOUT
# =========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# =========================================================
# ERROR HANDLERS
# =========================================================

@app.errorhandler(413)
def too_large(_):

    return render_template(
        "error.html",
        title="File too large",
        message="Maximum upload size is 10 MB."
    ), 413


@app.errorhandler(404)
def not_found(_):

    return render_template(
        "error.html",
        title="Page not found",
        message="The page you requested does not exist."
    ), 404


@app.errorhandler(500)
def server_error(_):

    return render_template(
        "error.html",
        title="Something went wrong",
        message="Please try again."
    ), 500


# =========================================================
# RUN
# =========================================================

if __name__ == "__main__":

    app.run(
        debug=os.getenv(
            "FLASK_DEBUG",
            "false"
        ).lower() == "true"
    )